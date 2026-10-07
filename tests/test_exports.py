"""Exports follow the profile; the DXMon bridge matches DXMon's own loader."""

import csv
import io

import pytest
from fastapi.testclient import TestClient

from app.categories import CW, MIXED
from app.main import create_app
from conftest import render_paste


@pytest.fixture
def client(tmp_path, reference, credits):
    app = create_app(db_path=tmp_path / "e.db", reference=reference, clublog_fetch=lambda: b"{}")
    with TestClient(app) as c:
        p = c.post("/api/import/preview", data={"text": render_paste(reference, credits)}).json()
        c.post(f"/api/import/{p['id']}/save")
        yield c


def _rows(content):
    return list(csv.DictReader(io.StringIO(content.decode("utf-8"))))


def test_nothing_to_export_before_import(tmp_path, reference):
    app = create_app(db_path=tmp_path / "x.db", reference=reference)
    with TestClient(app) as c:
        assert c.get("/export/no_confirms.csv").status_code == 409


def test_no_confirms_matches_dxmon_loader(client, reference, credits):
    r = client.get("/export/no_confirms.csv")
    assert r.status_code == 200 and 'filename="no_confirms.csv"' in r.headers["content-disposition"]
    raw = r.content
    assert not raw.startswith(b"\xef\xbb\xbf")          # no BOM: DXMon reads plain UTF-8
    assert raw.splitlines()[0] == b'Entity,Prefix'
    # Exactly DXMon's loader (adxo_service._load_no_confirms):
    rows = [{"entity": row["Entity"].strip(), "prefix": row.get("Prefix", "").strip()}
            for row in csv.DictReader(io.StringIO(raw.decode("utf-8"))) if row.get("Entity", "").strip()]
    missing = {e.lotw_name for e in reference.current if MIXED not in credits.get(e.dxcc, ())}
    assert {r["entity"] for r in rows} == missing and len(rows) == len(missing)
    spratly = [r for r in rows if r["entity"] == "SPRATLY ISLANDS"]
    if spratly:                                          # blank prefix survives
        assert spratly[0]["prefix"] == ""


def test_missing_entities_csv_carries_marks(client, reference, credits):
    new = next(e.dxcc for e in reference.current if e.dxcc not in credits)
    client.post("/api/marks", json={"dxcc": new, "category": MIXED, "state": "awaiting", "note": "QSL in"})
    client.post("/api/paper", json={"dxcc": new, "band": "20M", "mode": CW})
    rows = _rows(client.get("/export/missing_entities.csv").content)
    row = next(r for r in rows if int(r["DXCC"]) == new)
    assert row["Mark"] == "Awaiting" and row["Mark note"] == "QSL in" and row["Paper card"] == "In hand"


def test_missing_slots_follow_profile(client, reference, credits):
    rows = _rows(client.get("/export/missing_slots.csv").content)
    assert "2 m" not in rows[0] and "Satellite" not in rows[0] and "20 m" in rows[0]
    d = next(d for d, c in credits.items() if MIXED in c and CW not in c)
    client.post("/api/marks", json={"dxcc": d, "category": CW, "state": "wont_submit"})
    rows = _rows(client.get("/export/missing_slots.csv").content)
    assert next(r for r in rows if int(r["DXCC"]) == d)["CW"] == "Won't submit"
    client.put("/api/settings/profile", json={"bands": ["20M", "2M"], "modes": ["CW", "SAT"]})
    rows = _rows(client.get("/export/missing_slots.csv").content)
    assert "2 m" in rows[0] and "Satellite" in rows[0] and "40 m" not in rows[0]


def test_workbook_sheets(client):
    from openpyxl import load_workbook
    r = client.get("/export/workbook.xlsx")
    assert r.status_code == 200 and "dxcc_analysis_" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Missing Entities", "Missing Slots", "By Band",
                             "Matrix", "Pending Marks", "Paper QSLs"]
    assert wb["Summary"]["B2"].value == "N0CALL"
    assert wb["Matrix"].max_row == 341                    # header + 340 current entities
