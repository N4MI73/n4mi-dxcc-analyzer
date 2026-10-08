"""Tests for the pre-deployment review findings (2026-10-08)."""

import csv
import io

import pytest
from fastapi.testclient import TestClient

from app import db
from app.categories import CW, MIXED
from app.main import create_app
from app.reference import Reference
from conftest import render_paste


def _client(path, reference, **kw):
    return TestClient(create_app(db_path=path, reference=reference, clublog_fetch=lambda: b"{}", **kw))


def _import(c, text):
    p = c.post("/api/import/preview", data={"text": text}).json()
    assert c.post(f"/api/import/{p['id']}/save").status_code == 200
    return p


# 1. A damaged database or a missing data folder is reported, not crashed on.
def test_corrupt_database_is_reported(tmp_path, reference):
    bad = tmp_path / "dxcc.db"
    bad.write_bytes(b"this is not a sqlite database" * 100)
    with _client(bad, reference) as c:
        assert c.get("/healthz").json() == {"status": "ok"}
        st = c.get("/api/status").json()
        assert st["ok"] is False and "Database problem" in st["error"]
        r = c.get("/api/view/entities")
        assert r.status_code == 503 and "Database problem" in r.json()["detail"]


def test_missing_data_folder_is_reported_not_created(tmp_path, reference, monkeypatch):
    missing = tmp_path / "not-mounted"
    monkeypatch.setenv("DATA_DIR", str(missing))
    with TestClient(create_app(reference=reference)) as c:
        st = c.get("/api/status").json()
        assert st["ok"] is False and "not found" in st["error"]
    assert not missing.exists()


# 2. User-typed text never becomes a spreadsheet formula.
def test_formula_text_is_neutralised(tmp_path, reference, credits):
    from openpyxl import load_workbook
    with _client(tmp_path / "f.db", reference) as c:
        _import(c, render_paste(reference, credits))
        new = next(e.dxcc for e in reference.current if e.dxcc not in credits)
        c.post("/api/marks", json={"dxcc": new, "category": MIXED, "state": "awaiting",
                                   "note": '=HYPERLINK("http://x","click")'})
        c.post("/api/paper", json={"dxcc": new, "band": "20M", "mode": CW, "call": "=1+1",
                                   "note": "+44 bureau"})
        wb = load_workbook(io.BytesIO(c.get("/export/workbook.xlsx").content))
        for ws in wb:
            for row in ws.iter_rows():
                for cell in row:
                    assert cell.data_type != "f", (ws.title, cell.coordinate, cell.value)
        assert wb["Pending Marks"]["F2"].value == '=HYPERLINK("http://x","click")'   # as typed, as text
        rows = list(csv.DictReader(io.StringIO(c.get("/export/missing_entities.csv").text)))
        assert next(r for r in rows if int(r["DXCC"]) == new)["Mark note"].startswith("'=")


# 3. Data the entity table doesn't know is flagged and never crashes a page.
def test_reference_mismatch_is_reported(tmp_path, reference, credits):
    path = tmp_path / "m.db"
    gone = next(d for d, c in credits.items() if MIXED in c and CW not in c)
    with _client(path, reference) as c:
        _import(c, render_paste(reference, credits))
        c.post("/api/marks", json={"dxcc": gone, "category": CW, "state": "awaiting"})
    smaller = Reference([e for e in reference.entities if e.dxcc != gone])
    with _client(path, smaller) as c:
        st = c.get("/api/status").json()
        assert st["ok"] is False and gone in st["reference_problems"]
        marks = c.get("/api/marks").json()["marks"]
        assert marks[0]["name"] == f"Unknown DXCC {gone}"
        assert c.get("/export/workbook.xlsx").status_code == 200
        for v in ("entities", "slots", "matrix"):
            assert c.get(f"/api/view/{v}").status_code == 200


# 4. Promotion is atomic and checks state inside the transaction.
def test_promotion_refuses_stale_or_missing_snapshots(tmp_path, reference, credits):
    from app.validate import validate
    from app.parser import parse_text
    conn = db.connect(tmp_path / "c.db")
    db.init(conn)
    res = validate(parse_text(render_paste(reference, credits)), reference, None)
    first = db.store_preview(conn, res, "paste", held=False, based_on=None)
    db.make_current(conn, first, allowed=db.PENDING, expect_current=None)
    # A preview validated before `first` was saved (based_on=None) must be refused.
    stale = db.store_preview(conn, res, "paste", held=False, based_on=None)
    with pytest.raises(db.NotPromotable):
        db.make_current(conn, stale, allowed=db.PENDING, expect_current=None)
    # A missing or discarded snapshot can't be promoted, and the current one stays.
    with pytest.raises(db.NotPromotable):
        db.make_current(conn, 9999, allowed=db.PENDING)
    db.discard(conn, stale)
    with pytest.raises(db.NotPromotable):
        db.make_current(conn, stale, allowed=db.PENDING)
    assert db.current_snapshot(conn)["id"] == first
    conn.close()


# 5. Unreadable uploads get a clear message, not a server error.
def test_unreadable_upload(tmp_path, reference):
    with _client(tmp_path / "u.db", reference) as c:
        for name in ("m.csv", "m.txt"):
            r = c.post("/api/import/preview", files={"file": (name, b"\x81\x8d\x8f\x90\x9d")})
            assert r.status_code == 422 and "readable" in r.json()["detail"]


# 6. The history shows which marks an import cleared.
def test_history_lists_cleared_marks(tmp_path, reference, credits):
    with _client(tmp_path / "h.db", reference) as c:
        _import(c, render_paste(reference, credits))
        need = next(d for d, cr in credits.items() if MIXED in cr and CW not in cr)
        c.post("/api/marks", json={"dxcc": need, "category": CW, "state": "awaiting"})
        newer = dict(credits)
        newer[need] = credits[need] | {CW}
        second = _import(c, render_paste(reference, newer))
        snaps = {s["id"]: s for s in c.get("/api/snapshots").json()["snapshots"]}
        assert snaps[second["id"]]["cleared_marks"] == [
            {"name": reference.by_dxcc[need].lotw_name, "category": CW, "label": "CW"}]


# 7. Input sizes are limited.
def test_input_limits(tmp_path, reference, credits):
    with _client(tmp_path / "l.db", reference) as c:
        _import(c, render_paste(reference, credits))
        need = next(d for d, cr in credits.items() if MIXED in cr and CW not in cr)
        r = c.post("/api/marks", json={"dxcc": need, "category": CW, "state": "awaiting", "note": "x" * 501})
        assert r.status_code == 422
        # Refused either by Starlette's form-size limit (400) or by the app's own (413).
        r = c.post("/api/import/preview", data={"text": "x" * 2_000_001})
        assert r.status_code in (400, 413)
