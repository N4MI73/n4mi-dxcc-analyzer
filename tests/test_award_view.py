"""D51 award view: Band Slots, the Full Matrix and exports count never-credited
entities too, so each band/mode count equals LoTW's (current entities minus
credits in that category). Data is synthetic."""

import csv
import io

import pytest
from fastapi.testclient import TestClient

from app import analysis
from app.categories import DIGITAL, MIXED, SAT, Profile
from app.main import create_app
from conftest import render_paste


@pytest.fixture
def client(tmp_path, reference, credits):
    app = create_app(db_path=tmp_path / "a.db", reference=reference, clublog_fetch=lambda: b"{}")
    with TestClient(app) as c:
        p = c.post("/api/import/preview", data={"text": render_paste(reference, credits)}).json()
        c.post(f"/api/import/{p['id']}/save")
        yield c


def _never(reference, credits):
    return [e.dxcc for e in reference.current if MIXED not in credits.get(e.dxcc, ())]


def test_include_new_adds_every_slot_for_never_credited(reference, credits):
    p = Profile()
    confirmed = analysis.needed_slots(credits, reference, p)
    every = analysis.needed_slots(credits, reference, p, include_new=True)
    never = _never(reference, credits)
    assert never and all(d not in confirmed for d in never)
    slots = [c for c in p.slot_categories if c != SAT]
    assert all(every[d] == slots for d in never)
    assert {d: v for d, v in every.items() if d not in never} == confirmed


def test_summary_splits_confirmed_and_new(reference, credits):
    p = Profile()
    s = analysis.summarize(credits, reference, p)
    slots = len([c for c in p.slot_categories if c != SAT])
    assert s.slots_missing_new == len(_never(reference, credits)) * slots
    assert s.slots_missing == sum(len(v) for v in analysis.needed_slots(credits, reference, p).values())


def test_slot_counts_match_lotw_award_counts(client, reference, credits):
    tot = analysis.totals(credits)
    current = len(reference.current)
    never = set(_never(reference, credits))
    for c in client.get("/api/view/slots").json()["categories"]:
        assert len(c["entities"]) == current - tot[c["category"]], c["category"]
        assert c["new_count"] == len(never)
        assert c["confirmed_count"] + c["new_count"] == len(c["entities"])
        assert {e["dxcc"] for e in c["entities"] if e["new"]} == never


def test_never_credited_band_marks_show_everywhere(client, reference, credits):
    d = _never(reference, credits)[0]
    for cat in (MIXED, "20M", DIGITAL):
        r = client.post("/api/marks", json={"dxcc": d, "category": cat, "state": "awaiting",
                                            "note": "one QSO"})
        assert r.status_code == 200, r.text
    dig = next(c for c in client.get("/api/view/slots").json()["categories"] if c["category"] == DIGITAL)
    row = next(e for e in dig["entities"] if e["dxcc"] == d)
    assert row["new"] and row["marks"][DIGITAL]["state"] == "awaiting"
    m = next(e for e in client.get("/api/view/matrix").json()["entities"] if e["dxcc"] == d)
    assert m["new"] and set(m["marks"]) == {MIXED, "20M", DIGITAL}
    ent = next(e for e in client.get("/api/view/entities").json()["entities"] if e["dxcc"] == d)
    assert set(ent["marks"]) == {MIXED, "20M", DIGITAL}


def test_missing_slots_csv_flags_never_confirmed(client, reference, credits):
    rows = list(csv.DictReader(io.StringIO(client.get("/export/missing_slots.csv").content.decode())))
    never = set(_never(reference, credits))
    assert {int(r["DXCC"]) for r in rows if r["Never confirmed"] == "Yes"} == never
    assert all(r["Never confirmed"] == "" for r in rows if int(r["DXCC"]) not in never)
    # Column order before the new one is unchanged.
    assert list(rows[0])[:4] == ["DXCC", "Prefix", "Entity", "Slots needed"]
    assert list(rows[0])[-1] == "Never confirmed"


def test_workbook_by_band_award_counts(client, reference, credits):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get("/export/workbook.xlsx").content))
    ws = wb["By Band"]
    head = [c.value for c in ws[1]]
    assert head[:4] == ["Band or mode", "Entities needed", "On credited entities", "Never confirmed"]
    tot = analysis.totals(credits)
    dig = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "Digital")
    assert dig[1] == len(reference.current) - tot[DIGITAL]
    assert dig[1] == dig[2] + dig[3]


def test_rows_carry_credits_for_the_mark_menu(client, reference, credits):
    """The menu greys out credited slots, so every list row says what is credited."""
    for c in client.get("/api/view/slots").json()["categories"]:
        for e in c["entities"]:
            assert set(e["credited"]) == set(credits.get(e["dxcc"], ()))
    for e in client.get("/api/view/entities").json()["entities"]:
        assert MIXED not in e["credited"]
